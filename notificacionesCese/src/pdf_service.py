import os
import shutil
import subprocess
import unicodedata
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
import qrcode
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from config import (
    ACTA_INPUT_FIELDS,
    ACTA_PLACEHOLDER_MAP,
    ACTA_QR_IMAGE_WIDTH_INCHES,
    ACTA_QR_MARKER_TEXT,
    ACTA_TEMPLATE_FIELDS,
    ANEXO_QR_FILENAMES,
    RELACION_ANEXOS_OUTPUT_DOCX_FILENAME,
    RELACION_ANEXOS_OUTPUT_FILENAME,
    RELACION_ANEXOS_QR_IMAGE_WIDTH_INCHES,
    RELACION_ANEXOS_QR_OUTPUT_FILENAME,
)


def generate_qr_image(link: str, output_folder: Path | str, output_filename: str) -> str:
    if not link:
        raise ValueError("A non-empty link is required to generate a QR.")

    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=4,
    )
    qr.add_data(link)
    qr.make(fit=True)

    image = qr.make_image(fill_color="black", back_color="white")

    output_path = output_folder / output_filename
    image.save(str(output_path))
    return str(output_path)


def generate_anexo_qrs(validation_result: dict, output_folder: Path | str) -> dict:
    if not validation_result or "links" not in validation_result:
        raise ValueError("validation_result must contain a 'links' dictionary.")

    qr_files = {}

    for alias, filename in ANEXO_QR_FILENAMES.items():
        link = validation_result["links"].get(alias)
        if not link:
            raise ValueError(f"Missing link for '{alias}' in validation_result['links'].")

        qr_files[alias] = generate_qr_image(link, output_folder, filename)

    return qr_files



def _find_relacion_table(document: Document):
    for table in document.tables:
        if len(table.rows) < 2 or len(table.columns) < 3:
            continue

        first_row_text = [cell.text.strip().upper() for cell in table.rows[0].cells[:3]]
        if first_row_text == ["ANEXO A", "ANEXO B", "ANEXO C"]:
            return table

    raise ValueError(
        "Could not find the annex table with headers 'ANEXO A', 'ANEXO B' and 'ANEXO C' "
        "inside the relacion de anexos template."
    )


def generate_relacion_anexos_docx(
    qr_files: dict,
    output_folder: Path | str,
    template_docx_path: Path | str,
    output_filename: str = RELACION_ANEXOS_OUTPUT_DOCX_FILENAME,
) -> str:
    required_aliases = ("anexoA", "anexoB", "anexoC")

    missing_aliases = [alias for alias in required_aliases if alias not in qr_files]
    if missing_aliases:
        raise ValueError(
            f"qr_files is missing required keys: {', '.join(missing_aliases)}"
        )

    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    template_docx_path = Path(template_docx_path)
    if not template_docx_path.exists():
        raise FileNotFoundError(f"Template DOCX not found: {template_docx_path}")

    for alias in required_aliases:
        qr_path = Path(qr_files[alias])
        if not qr_path.exists():
            raise FileNotFoundError(f"QR image not found for '{alias}': {qr_path}")

    document = Document(str(template_docx_path))
    relacion_table = _find_relacion_table(document)

    aliases_in_order = ("anexoA", "anexoB", "anexoC")
    target_row = relacion_table.rows[1]

    for cell_index, alias in enumerate(aliases_in_order):
        cell = target_row.cells[cell_index]
        cell.text = ""

        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()
        run.add_picture(
            str(qr_files[alias]),
            width=Inches(RELACION_ANEXOS_QR_IMAGE_WIDTH_INCHES),
        )

    output_path = output_folder / output_filename
    document.save(str(output_path))
    return str(output_path)


def generate_relacion_anexos_qr(
    relacion_anexos_link: str,
    output_folder: Path | str,
    output_filename: str = RELACION_ANEXOS_QR_OUTPUT_FILENAME,
) -> str:
    return generate_qr_image(relacion_anexos_link, output_folder, output_filename)


_MONTH_NAMES = {
    1: "enero",
    2: "febrero",
    3: "marzo",
    4: "abril",
    5: "mayo",
    6: "junio",
    7: "julio",
    8: "agosto",
    9: "septiembre",
    10: "octubre",
    11: "noviembre",
    12: "diciembre",
}

_SPECIAL_NUMBERS = {
    0: "cero",
    1: "uno",
    2: "dos",
    3: "tres",
    4: "cuatro",
    5: "cinco",
    6: "seis",
    7: "siete",
    8: "ocho",
    9: "nueve",
    10: "diez",
    11: "once",
    12: "doce",
    13: "trece",
    14: "catorce",
    15: "quince",
    16: "dieciséis",
    17: "diecisiete",
    18: "dieciocho",
    19: "diecinueve",
    20: "veinte",
    21: "veintiuno",
    22: "veintidós",
    23: "veintitrés",
    24: "veinticuatro",
    25: "veinticinco",
    26: "veintiséis",
    27: "veintisiete",
    28: "veintiocho",
    29: "veintinueve",
}

_TENS = {
    30: "treinta",
    40: "cuarenta",
    50: "cincuenta",
    60: "sesenta",
    70: "setenta",
    80: "ochenta",
    90: "noventa",
}

_HUNDREDS = {
    100: "cien",
    200: "doscientos",
    300: "trescientos",
    400: "cuatrocientos",
    500: "quinientos",
    600: "seiscientos",
    700: "setecientos",
    800: "ochocientos",
    900: "novecientos",
}

_FIELD_NAME_ALIASES = {
    "numero de acta": "numero de acta",
    "fecha de acta": "fecha de acta",
    "fecha de diligencia": "fecha de diligencia",
    "hora inicio diligencia": "hora inicio diligencia",
    "direccion del restaurante": "dirección del restaurante",
    "direcci?n del restaurante": "dirección del restaurante",
    "nombre / denominacion del restaurante": "nombre / denominación del restaurante",
    "nombre / denominaci?n del restaurante": "nombre / denominación del restaurante",
    "fecha de entrega": "fecha de entrega",
    "sexo de quien recibio la carta": "sexo de quien recibió la carta",
    "sexo de quien recibi? la carta": "sexo de quien recibió la carta",
    "nombre de quien recibio la carta": "nombre de quien recibió la carta",
    "nombre de quien recibi? la carta": "nombre de quien recibió la carta",
    "clausula_firma": "clausula_firma",
    "hora fin diligencia": "hora fin diligencia",
    "numero de instrumento": "numero de instrumento",
    "fecha del instrumento": "fecha del instrumento",
}


def _canonicalize_field_name(field_name: str) -> str:
    normalized = unicodedata.normalize("NFKD", str(field_name).strip().lower())
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    return normalized


_EXPECTED_FIELDS_BY_CANONICAL = {
    _canonicalize_field_name(field_name): field_name for field_name in ACTA_INPUT_FIELDS
}


def _normalize_text_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number_to_words_under_1000(number: int) -> str:
    if number < 30:
        return _SPECIAL_NUMBERS[number]
    if number < 100:
        tens = (number // 10) * 10
        units = number % 10
        if units == 0:
            return _TENS[tens]
        return f"{_TENS[tens]} y {_SPECIAL_NUMBERS[units]}"
    if number == 100:
        return "cien"

    hundreds = (number // 100) * 100
    remainder = number % 100
    prefix = "ciento" if number < 200 else _HUNDREDS[hundreds]
    if remainder == 0:
        return prefix
    return f"{prefix} {_number_to_words_under_1000(remainder)}"


def _apocopate_uno(text: str) -> str:
    if text.endswith("veintiuno"):
        return text[:-9] + "veintiún"
    if text.endswith(" y uno"):
        return text[:-6] + " y un"
    if text.endswith(" uno"):
        return text[:-4] + " un"
    if text == "uno":
        return "un"
    return text


def number_to_spanish_words(number: int, *, apocopate: bool = False) -> str:
    if number < 0:
        raise ValueError("Only positive numbers are supported.")

    if number < 1000:
        result = _number_to_words_under_1000(number)
        return _apocopate_uno(result) if apocopate else result

    if number < 1_000_000:
        thousands = number // 1000
        remainder = number % 1000
        if thousands == 1:
            prefix = "mil"
        else:
            prefix = f"{number_to_spanish_words(thousands, apocopate=True)} mil"
        if remainder == 0:
            return prefix
        result = f"{prefix} {_number_to_words_under_1000(remainder)}"
        return _apocopate_uno(result) if apocopate else result

    raise ValueError("Only numbers below one million are supported.")


def _parse_int_field(value, field_name: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{field_name} must be a number, not a boolean value.")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"{field_name} must be a whole number.")
        return int(value)

    text = _normalize_text_value(value).replace(",", "").replace(" ", "")
    if not text.isdigit():
        raise ValueError(f"{field_name} must contain only digits.")
    return int(text)


def _parse_date_field(value, field_name: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = _normalize_text_value(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    raise ValueError(
        f"{field_name} must be a valid date like 25/03/2026 or 2026-03-25."
    )


def _parse_time_field(value, field_name: str) -> time:
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)

    text = _normalize_text_value(value)
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p"):
        try:
            return datetime.strptime(text, fmt).time().replace(second=0, microsecond=0)
        except ValueError:
            continue

    raise ValueError(f"{field_name} must be a valid time like 12:50 or 13:02.")


def date_to_spanish_words(value: date, *, uppercase: bool = False) -> str:
    text = (
        f"{number_to_spanish_words(value.day)} de "
        f"{_MONTH_NAMES[value.month]} de "
        f"{number_to_spanish_words(value.year)}"
    )
    return text.upper() if uppercase else text


def time_to_spanish_words(value: time, *, uppercase: bool = False) -> str:
    hour_word = "una" if value.hour == 1 else number_to_spanish_words(value.hour)
    hour_label = "hora" if value.hour == 1 else "horas"
    minute_label = "minuto" if value.minute == 1 else "minutos"
    text = (
        f"{hour_word} {hour_label} con "
        f"{number_to_spanish_words(value.minute)} {minute_label}"
    )
    return text.upper() if uppercase else text


def build_acta_template_data(raw_data: dict) -> dict:
    numero_acta = _parse_int_field(raw_data["numero de acta"], "numero de acta")
    fecha_acta = _parse_date_field(raw_data["fecha de acta"], "fecha de acta")
    fecha_diligencia = _parse_date_field(
        raw_data["fecha de diligencia"], "fecha de diligencia"
    )
    hora_inicio = _parse_time_field(
        raw_data["hora inicio diligencia"], "hora inicio diligencia"
    )
    hora_fin = _parse_time_field(raw_data["hora fin diligencia"], "hora fin diligencia")
    instrumento_numero = _parse_int_field(
        raw_data["numero de instrumento"], "numero de instrumento"
    )
    instrumento_fecha = _parse_date_field(
        raw_data["fecha del instrumento"], "fecha del instrumento"
    )
    fecha_entrega = _parse_date_field(raw_data["fecha de entrega"], "fecha de entrega")

    result = {
        "acta_numero": str(numero_acta),
        "acta_numero_letra": number_to_spanish_words(numero_acta, apocopate=True).upper(),
        "fecha_acta_num": fecha_acta.strftime("%d/%m/%Y"),
        "fecha_acta_letra": date_to_spanish_words(fecha_acta, uppercase=True),
        "fecha_diligencia_letra": date_to_spanish_words(fecha_diligencia),
        "hora_inicio_letra": time_to_spanish_words(hora_inicio, uppercase=True),
        "direccion_restaurante": raw_data["dirección del restaurante"],
        "nombre_restaurante": raw_data["nombre / denominación del restaurante"],
        "fecha_entrega_letra": date_to_spanish_words(fecha_entrega),
        "sexo_recibe": raw_data["sexo de quien recibió la carta"],
        "nombre_recibe": raw_data["nombre de quien recibió la carta"],
        "clausula_firma": raw_data["clausula_firma"],
        "hora_fin_letra": time_to_spanish_words(hora_fin, uppercase=True),
        "instrumento_numero_letra": number_to_spanish_words(
            instrumento_numero,
            apocopate=True,
        ),
        "instrumento_fecha_letra": date_to_spanish_words(instrumento_fecha),
    }

    empty_fields = [field for field in ACTA_TEMPLATE_FIELDS if not result[field].strip()]
    if empty_fields:
        raise ValueError(
            "The following generated acta fields are empty: " + ", ".join(empty_fields)
        )

    return result


def read_acta_data_xlsx(local_xlsx_path: str | Path) -> dict:
    local_xlsx_path = Path(local_xlsx_path)
    if not local_xlsx_path.exists():
        raise FileNotFoundError(f"acta_data.xlsx not found: {local_xlsx_path}")

    workbook = openpyxl.load_workbook(local_xlsx_path, data_only=True)
    worksheet = workbook.active

    result = {}

    for row in worksheet.iter_rows(min_row=6, values_only=True):
        field_name = row[0]
        field_value = row[1]

        if field_name is None:
            continue

        raw_field_name = str(field_name).strip().lower()
        aliased_field_name = _FIELD_NAME_ALIASES.get(raw_field_name, raw_field_name)
        normalized_field_name = _EXPECTED_FIELDS_BY_CANONICAL.get(
            _canonicalize_field_name(aliased_field_name),
            aliased_field_name,
        )
        normalized_value = _normalize_text_value(field_value)
        result[normalized_field_name] = normalized_value

    missing_fields = [field for field in ACTA_INPUT_FIELDS if field not in result]
    if missing_fields:
        raise ValueError(
            "acta_data.xlsx is missing required fields: " + ", ".join(missing_fields)
        )

    empty_fields = [field for field in ACTA_INPUT_FIELDS if not result.get(field, "").strip()]
    if empty_fields:
        raise ValueError(
            "The following acta fields are empty in acta_data.xlsx: " + ", ".join(empty_fields)
        )

    return build_acta_template_data({field: result[field] for field in ACTA_INPUT_FIELDS})


def _iter_all_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph

    for section in document.sections:
        for container in (section.header, section.first_page_header, section.even_page_header):
            for paragraph in container.paragraphs:
                yield paragraph
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            yield paragraph

        for container in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in container.paragraphs:
                yield paragraph
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            yield paragraph


def _clear_run_marker_formatting(run):
    run.font.highlight_color = None
    run_element = run._element
    rpr = run_element.rPr
    if rpr is None:
        return

    shading = rpr.find(qn("w:shd"))
    if shading is not None:
        rpr.remove(shading)


def _replace_placeholder_in_paragraph(paragraph, placeholder: str, replacement: str):
    if not paragraph.runs:
        return

    full_text = "".join(run.text for run in paragraph.runs)
    if placeholder not in full_text:
        return

    while placeholder in full_text:
        start = full_text.index(placeholder)
        end = start + len(placeholder)

        current_index = 0
        start_run_index = None
        start_offset = None
        end_run_index = None
        end_offset = None

        for run_index, run in enumerate(paragraph.runs):
            run_text = run.text
            run_start = current_index
            run_end = current_index + len(run_text)

            if start_run_index is None and start < run_end:
                start_run_index = run_index
                start_offset = start - run_start

            if end <= run_end:
                end_run_index = run_index
                end_offset = end - run_start
                break

            current_index = run_end

        if start_run_index is None or end_run_index is None:
            break

        start_run = paragraph.runs[start_run_index]
        end_run = paragraph.runs[end_run_index]

        prefix = start_run.text[:start_offset]
        suffix = end_run.text[end_offset:]

        start_run.text = prefix + replacement + suffix

        _clear_run_marker_formatting(start_run)

        for run_index in range(start_run_index + 1, end_run_index + 1):
            paragraph.runs[run_index].text = ""
            _clear_run_marker_formatting(paragraph.runs[run_index])

        full_text = "".join(run.text for run in paragraph.runs)


def replace_placeholders_in_docx(document: Document, replacements: dict):
    for paragraph in _iter_all_paragraphs(document):
        for placeholder, replacement in replacements.items():
            _replace_placeholder_in_paragraph(paragraph, placeholder, replacement)

def insert_qr_into_acta_docx(document: Document, qr_image_path: str | Path):
    qr_image_path = Path(qr_image_path)
    if not qr_image_path.exists():
        raise FileNotFoundError(f"QR image not found: {qr_image_path}")

    for table in document.tables:
        for row in table.rows:
            for cell_index, cell in enumerate(row.cells):
                if ACTA_QR_MARKER_TEXT in cell.text:
                    # El marcador está en la celda derecha.
                    # Queremos reemplazar el recuadro/indicador de la izquierda.
                    target_cell = row.cells[cell_index - 1] if cell_index > 0 else cell

                    # Vaciar la celda del recuadro
                    target_cell.text = ""

                    # Vaciar también la celda del marcador
                    cell.text = ""

                    paragraph = target_cell.paragraphs[0]
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    run = paragraph.add_run()
                    run.add_picture(
                        str(qr_image_path),
                        width=Inches(ACTA_QR_IMAGE_WIDTH_INCHES),
                    )
                    return

    raise ValueError("Could not find the QR marker cell inside the acta template.")

def generate_acta_docx(
    acta_data: dict,
    relacion_anexos_qr_path: str | Path,
    output_folder: Path | str,
    template_docx_path: Path | str,
    output_filename: str = "ACTA.docx",
) -> str:
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    template_docx_path = Path(template_docx_path)
    if not template_docx_path.exists():
        raise FileNotFoundError(f"Acta template DOCX not found: {template_docx_path}")

    document = Document(str(template_docx_path))

    replacements = {
        placeholder: acta_data[field_name]
        for placeholder, field_name in ACTA_PLACEHOLDER_MAP.items()
    }

    replace_placeholders_in_docx(document, replacements)
    insert_qr_into_acta_docx(document, relacion_anexos_qr_path)

    output_path = output_folder / output_filename
    document.save(str(output_path))
    return str(output_path)


def convert_docx_to_pdf(
    docx_path: str | Path,
    output_folder: Path | str,
    output_filename: str = "ACTA.pdf",
) -> str:
    docx_path = Path(docx_path).resolve()
    if not docx_path.exists():
        raise FileNotFoundError(f"DOCX file not found: {docx_path}")

    output_folder = Path(output_folder).resolve()
    output_folder.mkdir(parents=True, exist_ok=True)

    output_pdf_path = output_folder / output_filename
    if output_pdf_path.exists():
        output_pdf_path.unlink()

    if os.name == "nt":
        try:
            from docx2pdf import convert as docx2pdf_convert

            docx2pdf_convert(str(docx_path), str(output_pdf_path))
            if output_pdf_path.exists():
                return str(output_pdf_path)
        except Exception:
            pass

    libreoffice_binary = shutil.which("soffice") or shutil.which("libreoffice")
    if libreoffice_binary:
        command = [
            libreoffice_binary,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_folder),
            str(docx_path),
        ]

        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            check=False,
        )

        converted_pdf_path = output_folder / f"{docx_path.stem}.pdf"

        if completed.returncode == 0 and converted_pdf_path.exists():
            if converted_pdf_path != output_pdf_path:
                if output_pdf_path.exists():
                    output_pdf_path.unlink()
                converted_pdf_path.replace(output_pdf_path)
            return str(output_pdf_path)

        stderr = (completed.stderr or "").strip()
        stdout = (completed.stdout or "").strip()
        details = stderr or stdout or "LibreOffice failed without additional output."
        raise RuntimeError(f"Could not convert DOCX to PDF with LibreOffice. {details}")

    raise RuntimeError(
        "Could not convert DOCX to PDF. On Windows, install 'docx2pdf' and Microsoft Word. "
        "On Windows/Linux/macOS, you can also install LibreOffice and make sure 'soffice' is available in PATH."
    )
