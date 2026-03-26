import os
import shutil
import subprocess
from pathlib import Path

import openpyxl
import qrcode
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from config import (
    ACTA_OPTIONAL_FIELDS,
    ACTA_PLACEHOLDER_MAP,
    ACTA_QR_IMAGE_WIDTH_INCHES,
    ACTA_QR_MARKER_TEXT,
    ACTA_REQUIRED_FIELDS,
    ANEXO_QR_FILENAMES,
    RELACION_ANEXOS_OUTPUT_DOCX_FILENAME,
    RELACION_ANEXOS_OUTPUT_FILENAME,
    RELACION_ANEXOS_QR_IMAGE_WIDTH_INCHES,
    RELACION_ANEXOS_QR_OUTPUT_FILENAME,
)


def generate_qr_image(link: str, output_folder: Path | str, output_filename: str) -> str:
    if not link:
        raise ValueError("A non-empty link is required to generate a QR.")

    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=4,
    )
    qr.add_data(link)
    qr.make(fit=True)

    image = qr.make_image(fill_color="black", back_color="white")

    output_path = output_folder / output_filename
    image.save(str(output_path))
    return str(output_path)


def generate_anexo_qrs(validation_result: dict, output_folder: Path | str) -> dict:
    if not validation_result or "links" not in validation_result:
        raise ValueError("validation_result must contain a 'links' dictionary.")

    qr_files = {}

    for alias, filename in ANEXO_QR_FILENAMES.items():
        link = validation_result["links"].get(alias)
        if not link:
            raise ValueError(f"Missing link for '{alias}' in validation_result['links'].")

        qr_files[alias] = generate_qr_image(link, output_folder, filename)

    return qr_files



def _find_relacion_table(document: Document):
    for table in document.tables:
        if len(table.rows) < 2 or len(table.columns) < 3:
            continue

        first_row_text = [cell.text.strip().upper() for cell in table.rows[0].cells[:3]]
        if first_row_text == ["ANEXO A", "ANEXO B", "ANEXO C"]:
            return table

    raise ValueError(
        "Could not find the annex table with headers 'ANEXO A', 'ANEXO B' and 'ANEXO C' "
        "inside the relacion de anexos template."
    )


def generate_relacion_anexos_docx(
    qr_files: dict,
    output_folder: Path | str,
    template_docx_path: Path | str,
    output_filename: str = RELACION_ANEXOS_OUTPUT_DOCX_FILENAME,
) -> str:
    required_aliases = ("anexoA", "anexoB", "anexoC")

    missing_aliases = [alias for alias in required_aliases if alias not in qr_files]
    if missing_aliases:
        raise ValueError(
            f"qr_files is missing required keys: {', '.join(missing_aliases)}"
        )

    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    template_docx_path = Path(template_docx_path)
    if not template_docx_path.exists():
        raise FileNotFoundError(f"Template DOCX not found: {template_docx_path}")

    for alias in required_aliases:
        qr_path = Path(qr_files[alias])
        if not qr_path.exists():
            raise FileNotFoundError(f"QR image not found for '{alias}': {qr_path}")

    document = Document(str(template_docx_path))
    relacion_table = _find_relacion_table(document)

    aliases_in_order = ("anexoA", "anexoB", "anexoC")
    target_row = relacion_table.rows[1]

    for cell_index, alias in enumerate(aliases_in_order):
        cell = target_row.cells[cell_index]
        cell.text = ""

        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()
        run.add_picture(
            str(qr_files[alias]),
            width=Inches(RELACION_ANEXOS_QR_IMAGE_WIDTH_INCHES),
        )

    output_path = output_folder / output_filename
    document.save(str(output_path))
    return str(output_path)


def generate_relacion_anexos_qr(
    relacion_anexos_link: str,
    output_folder: Path | str,
    output_filename: str = RELACION_ANEXOS_QR_OUTPUT_FILENAME,
) -> str:
    return generate_qr_image(relacion_anexos_link, output_folder, output_filename)


def read_acta_data_xlsx(local_xlsx_path: str | Path) -> dict:
    local_xlsx_path = Path(local_xlsx_path)
    if not local_xlsx_path.exists():
        raise FileNotFoundError(f"acta_data.xlsx not found: {local_xlsx_path}")

    workbook = openpyxl.load_workbook(local_xlsx_path, data_only=True)
    worksheet = workbook.active

    result = {}

    for row in worksheet.iter_rows(min_row=6, values_only=True):
        field_name = row[0]
        field_value = row[1]

        if field_name is None:
            continue

        normalized_field_name = str(field_name).strip().lower()
        normalized_value = "" if field_value is None else str(field_value).strip()
        result[normalized_field_name] = normalized_value

    missing_fields = [field for field in ACTA_REQUIRED_FIELDS if field not in result]
    if missing_fields:
        raise ValueError(
            "acta_data.xlsx is missing required fields: " + ", ".join(missing_fields)
        )

    empty_fields = [field for field in ACTA_REQUIRED_FIELDS if not result.get(field, "").strip()]
    if empty_fields:
        raise ValueError(
            "The following acta fields are empty in acta_data.xlsx: " + ", ".join(empty_fields)
        )

    normalized_result = {field: result[field] for field in ACTA_REQUIRED_FIELDS}
    for field in ACTA_OPTIONAL_FIELDS:
        normalized_result[field] = result.get(field, "").strip()

    return normalized_result


def build_canal_tv_fragment(acta_data: dict) -> str:
    canal_tv = acta_data.get("canal tv", "").strip()
    if not canal_tv:
        return ""

    return f', del canal "{canal_tv}"'


def _iter_all_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _replace_placeholder_in_paragraph(paragraph, placeholder: str, replacement: str):
    if not paragraph.runs:
        return

    full_text = "".join(run.text for run in paragraph.runs)
    if placeholder not in full_text:
        return

    while placeholder in full_text:
        start = full_text.index(placeholder)
        end = start + len(placeholder)

        current_index = 0
        start_run_index = None
        start_offset = None
        end_run_index = None
        end_offset = None

        for run_index, run in enumerate(paragraph.runs):
            run_text = run.text
            run_start = current_index
            run_end = current_index + len(run_text)

            if start_run_index is None and start < run_end:
                start_run_index = run_index
                start_offset = start - run_start

            if end <= run_end:
                end_run_index = run_index
                end_offset = end - run_start
                break

            current_index = run_end

        if start_run_index is None or end_run_index is None:
            break

        start_run = paragraph.runs[start_run_index]
        end_run = paragraph.runs[end_run_index]

        prefix = start_run.text[:start_offset]
        suffix = end_run.text[end_offset:]

        start_run.text = prefix + replacement + suffix

        for run_index in range(start_run_index + 1, end_run_index + 1):
            paragraph.runs[run_index].text = ""

        # Quitar resaltado amarillo del texto reemplazado
        start_run.font.highlight_color = None

        full_text = "".join(run.text for run in paragraph.runs)


def replace_placeholders_in_docx(document: Document, replacements: dict):
    for paragraph in _iter_all_paragraphs(document):
        for placeholder, replacement in replacements.items():
            _replace_placeholder_in_paragraph(paragraph, placeholder, replacement)

def insert_qr_into_acta_docx(document: Document, qr_image_path: str | Path):
    qr_image_path = Path(qr_image_path)
    if not qr_image_path.exists():
        raise FileNotFoundError(f"QR image not found: {qr_image_path}")

    for table in document.tables:
        for row in table.rows:
            for cell_index, cell in enumerate(row.cells):
                if ACTA_QR_MARKER_TEXT in cell.text:
                    # El marcador está en la celda derecha.
                    # Queremos reemplazar el recuadro/indicador de la izquierda.
                    target_cell = row.cells[cell_index - 1] if cell_index > 0 else cell

                    # Vaciar la celda del recuadro
                    target_cell.text = ""

                    # Vaciar también la celda del marcador
                    cell.text = ""

                    paragraph = target_cell.paragraphs[0]
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    run = paragraph.add_run()
                    run.add_picture(
                        str(qr_image_path),
                        width=Inches(ACTA_QR_IMAGE_WIDTH_INCHES),
                    )
                    return

    raise ValueError("Could not find the QR marker cell inside the acta template.")

def generate_acta_docx(
    acta_data: dict,
    relacion_anexos_qr_path: str | Path,
    output_folder: Path | str,
    template_docx_path: Path | str,
    output_filename: str = "ACTA.docx",
) -> str:
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    template_docx_path = Path(template_docx_path)
    if not template_docx_path.exists():
        raise FileNotFoundError(f"Acta template DOCX not found: {template_docx_path}")

    document = Document(str(template_docx_path))

    replacements = {}
    for placeholder, field_name in ACTA_PLACEHOLDER_MAP.items():
        if field_name == "canal tv fragmento":
            replacements[placeholder] = build_canal_tv_fragment(acta_data)
            continue

        replacements[placeholder] = acta_data[field_name]

    replace_placeholders_in_docx(document, replacements)
    insert_qr_into_acta_docx(document, relacion_anexos_qr_path)

    output_path = output_folder / output_filename
    document.save(str(output_path))
    return str(output_path)


def convert_docx_to_pdf(
    docx_path: str | Path,
    output_folder: Path | str,
    output_filename: str = "ACTA.pdf",
) -> str:
    docx_path = Path(docx_path).resolve()
    if not docx_path.exists():
        raise FileNotFoundError(f"DOCX file not found: {docx_path}")

    output_folder = Path(output_folder).resolve()
    output_folder.mkdir(parents=True, exist_ok=True)

    output_pdf_path = output_folder / output_filename
    if output_pdf_path.exists():
        output_pdf_path.unlink()

    if os.name == "nt":
        try:
            from docx2pdf import convert as docx2pdf_convert

            docx2pdf_convert(str(docx_path), str(output_pdf_path))
            if output_pdf_path.exists():
                return str(output_pdf_path)
        except Exception:
            pass

    libreoffice_binary = shutil.which("soffice") or shutil.which("libreoffice")
    if libreoffice_binary:
        command = [
            libreoffice_binary,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_folder),
            str(docx_path),
        ]

        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            check=False,
        )

        converted_pdf_path = output_folder / f"{docx_path.stem}.pdf"

        if completed.returncode == 0 and converted_pdf_path.exists():
            if converted_pdf_path != output_pdf_path:
                if output_pdf_path.exists():
                    output_pdf_path.unlink()
                converted_pdf_path.replace(output_pdf_path)
            return str(output_pdf_path)

        stderr = (completed.stderr or "").strip()
        stdout = (completed.stdout or "").strip()
        details = stderr or stdout or "LibreOffice failed without additional output."
        raise RuntimeError(f"Could not convert DOCX to PDF with LibreOffice. {details}")

    raise RuntimeError(
        "Could not convert DOCX to PDF. On Windows, install 'docx2pdf' and Microsoft Word. "
        "On Windows/Linux/macOS, you can also install LibreOffice and make sure 'soffice' is available in PATH."
    )
